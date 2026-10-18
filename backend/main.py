
import io
import os
import re
from contextlib import contextmanager

import requests
from dotenv import load_dotenv
from fastapi import (
    FastAPI,
    HTTPException,
    Query,
    UploadFile,
    File,
    Form,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pypdf import PdfReader
from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
)
from xml.sax.saxutils import escape


load_dotenv()

app = FastAPI(title="OpportunityOS API")

default_origins = [
    "http://localhost:3000",
    "http://127.0.0.1:3000",
    "https://vx8gqsg2-3000.inc1.devtunnels.ms",
]
configured_origins = [
    origin.strip()
    for origin in os.getenv("CORS_ORIGINS", "").split(",")
    if origin.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=list(dict.fromkeys(default_origins + configured_origins)),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_FILE_SIZE = 5 * 1024 * 1024

SKILL_KEYWORDS = [
    "Python",
    "Java",
    "JavaScript",
    "TypeScript",
    "C++",
    "SQL",
    "HTML",
    "CSS",
    "React",
    "Next.js",
    "Node.js",
    "FastAPI",
    "Django",
    "Flask",
    "Machine Learning",
    "Deep Learning",
    "Artificial Intelligence",
    "Data Analysis",
    "Data Science",
    "Excel",
    "Power BI",
    "Tableau",
    "Statistics",
    "TensorFlow",
    "PyTorch",
    "Natural Language Processing",
    "NLP",
    "Generative AI",
    "Large Language Models",
    "LLM",
    "Git",
    "GitHub",
    "Docker",
    "AWS",
    "Azure",
    "Google Cloud",
    "REST API",
    "PostgreSQL",
    "MongoDB",
    "MySQL",
    "Communication",
    "Leadership",
    "Project Management",
    "Problem Solving",
    "Teamwork",
    "Data Structures",
    "Algorithms",
]


@app.get("/")
def home():
    return {
        "message": "OpportunityOS API is running!",
        "features": [
            "Live job search",
            "Resume analysis",
            "Tailored resume PDF generation",
        ],
    }


# --------------------------------------------------
# LIVE JOB SEARCH
# --------------------------------------------------

@app.get("/jobs")
def search_jobs(
    query: str = Query(
        default="software developer jobs in India",
        min_length=2,
        max_length=200,
    )
):
    api_key = os.getenv("RAPIDAPI_KEY")

    if not api_key:
        raise HTTPException(
            status_code=500,
            detail="RAPIDAPI_KEY is missing from backend/.env",
        )

    try:
        response = requests.get(
            "https://jsearch.p.rapidapi.com/search-v2",
            headers={
                "X-RapidAPI-Key": api_key,
                "X-RapidAPI-Host": "jsearch.p.rapidapi.com",
            },
            params={
                "query": query,
                "page": "1",
                "num_pages": "1",
                "country": "in",
            },
            timeout=30,
        )

        response.raise_for_status()
        result = response.json()

        if result.get("status") == "ERROR":
            raise HTTPException(
                status_code=502,
                detail="JSearch returned an API error.",
            )

        data = result.get("data", [])

        if isinstance(data, dict):
            jobs = data.get("jobs", [])
            cursor = data.get("cursor")
        elif isinstance(data, list):
            jobs = data
            cursor = None
        else:
            jobs = []
            cursor = None

        return {
            "jobs": jobs,
            "count": len(jobs),
            "cursor": cursor,
        }

    except requests.HTTPError as e:
        raise HTTPException(
            status_code=502,
            detail=(
                f"JSearch HTTP {e.response.status_code}: "
                f"{e.response.text[:300]}"
            ),
        ) from e

    except requests.RequestException as e:
        raise HTTPException(
            status_code=502,
            detail=f"Unable to connect to JSearch: {type(e).__name__}",
        ) from e


# --------------------------------------------------
# RESUME FILE READING
# --------------------------------------------------

async def read_resume(file: UploadFile) -> str:
    filename = (file.filename or "").lower()

    if not filename.endswith((".pdf", ".docx", ".txt")):
        raise HTTPException(
            status_code=400,
            detail="Upload a PDF, DOCX, or TXT resume.",
        )

    content = await file.read(MAX_FILE_SIZE + 1)

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="Resume must be 5 MB or smaller.",
        )

    try:
        if filename.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(content))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDFs are not supported.",
                )

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        elif filename.endswith(".docx"):
            document = Document(io.BytesIO(content))

            text = "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
            )

            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(
                        cell.text for cell in row.cells
                    )

        else:
            text = content.decode(
                "utf-8-sig",
                errors="replace",
            )

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read this resume. "
                "Check that the file is valid."
            ),
        ) from e

    text = text.strip()

    if len(text) < 30:
        raise HTTPException(
            status_code=400,
            detail=(
                "Very little text could be extracted. "
                "If this is a scanned PDF, OCR is not yet supported."
            ),
        )

    return text


def find_skills(text: str) -> list[str]:
    found = []
    normalized_text = text.lower()

    for skill in SKILL_KEYWORDS:
        pattern = (
            r"(?<!\w)"
            + re.escape(skill.lower())
            + r"(?!\w)"
        )

        if re.search(pattern, normalized_text):
            found.append(skill)

    return found


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


# --------------------------------------------------
# RESUME ANALYSIS
# --------------------------------------------------

@app.post("/resume/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    job_description: str = Form(default=""),
):
    resume_text = await read_resume(file)
    resume_skills = find_skills(resume_text)

    job_skills = (
        find_skills(job_description)
        if job_description.strip()
        else []
    )

    resume_skill_set = {
        skill.lower() for skill in resume_skills
    }

    matched_skills = [
        skill for skill in job_skills
        if skill.lower() in resume_skill_set
    ]

    missing_skills = [
        skill for skill in job_skills
        if skill.lower() not in resume_skill_set
    ]

    sections = [
        section
        for section in [
            "education",
            "experience",
            "projects",
            "skills",
            "certifications",
        ]
        if section in resume_text.lower()
    ]

    return {
        "filename": file.filename,
        "character_count": len(resume_text),
        "skills_found": resume_skills,
        "skills_count": len(resume_skills),
        "sections_detected": sections,
        "job_skills_found": job_skills,
        "matching_skills": matched_skills,
        "potential_skill_gaps": missing_skills,
        "resume_text": resume_text,
        "note": (
            "This is a keyword-based analysis, not an AI judgement. "
            "Skill gaps are suggestions to review, not proof of "
            "ineligibility."
        ),
    }


# --------------------------------------------------
# NON-AI RESUME PDF GENERATION
# --------------------------------------------------

@app.post("/resume/generate")
async def generate_resume(
    file: UploadFile = File(...),
    job_title: str = Form(default="Target Position"),
    job_description: str = Form(default=""),
):
    resume_text = await read_resume(file)

    job_title = clean_text(job_title) or "Target Position"
    job_description = clean_text(job_description)

    original_skills = find_skills(resume_text)
    job_skills = find_skills(job_description)

    original_skill_set = {
        skill.lower() for skill in original_skills
    }

    matching_skills = [
        skill for skill in job_skills
        if skill.lower() in original_skill_set
    ]

    missing_skills = [
        skill for skill in job_skills
        if skill.lower() not in original_skill_set
    ]

    buffer = io.BytesIO()

    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=f"OpportunityOS Resume - {job_title}",
        author="OpportunityOS",
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="ResumeTitleOS",
            parent=styles["Title"],
            fontSize=19,
            leading=24,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#172554"),
            spaceAfter=8,
        )
    )

    styles.add(
        ParagraphStyle(
            name="ResumeHeadingOS",
            parent=styles["Heading2"],
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#ea580c"),
            spaceBefore=12,
            spaceAfter=5,
            keepWithNext=True,
        )
    )

    styles.add(
        ParagraphStyle(
            name="ResumeBodyOS",
            parent=styles["BodyText"],
            fontSize=9,
            leading=13,
            spaceAfter=5,
            wordWrap="CJK",
        )
    )

    story = []

    def add_text(text, style="ResumeBodyOS"):
        text = clean_text(str(text))

        if text:
            story.append(
                Paragraph(
                    escape(text),
                    styles[style],
                )
            )

    story.append(
        Paragraph(
            "TAILORED RESUME",
            styles["ResumeTitleOS"],
        )
    )

    add_text(f"Target role: {job_title}")

    story.append(Spacer(1, 4 * mm))

    story.append(
        HRFlowable(
            width="100%",
            thickness=1,
            color=colors.HexColor("#ea580c"),
        )
    )

    if matching_skills:
        story.append(
            Paragraph(
                "Relevant Skills Identified",
                styles["ResumeHeadingOS"],
            )
        )
        add_text(", ".join(matching_skills))

    if missing_skills:
        story.append(
            Paragraph(
                "Skills to Review",
                styles["ResumeHeadingOS"],
            )
        )
        add_text(", ".join(missing_skills))
        add_text(
            "These job-description skills were not detected in "
            "your original resume. Only add them if you genuinely "
            "possess them."
        )

    story.append(
        Paragraph(
            "Original Resume Content",
            styles["ResumeHeadingOS"],
        )
    )

    headings = {
        "PROFESSIONAL SUMMARY",
        "SUMMARY",
        "EDUCATION",
        "TECHNICAL SKILLS",
        "SKILLS",
        "PROJECTS",
        "EXPERIENCE",
        "WORK EXPERIENCE",
        "CERTIFICATIONS & ACTIVITIES",
        "CERTIFICATIONS",
        "ACHIEVEMENTS",
    }

    for line in resume_text.splitlines():
        line = clean_text(line)

        if not line:
            continue

        normalized = line.strip().rstrip(":").upper()

        if normalized in headings:
            story.append(
                Paragraph(
                    escape(normalized),
                    styles["ResumeHeadingOS"],
                )
            )
        else:
            add_text(line)

    story.append(Spacer(1, 4 * mm))

    add_text(
        "Review your resume carefully before using it professionally. "
        "This version formats the original resume and identifies "
        "relevant skills; it does not rewrite your experience using AI."
    )

    pdf.build(story)
    buffer.seek(0)

    safe_title = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        job_title,
    ).strip("_")[:50] or "resume"

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="OpportunityOS_{safe_title}.pdf"'
        },
    )


# =========================
# PROFILE AGENT
# =========================

import sqlite3
import json
from pathlib import Path
from pydantic import BaseModel, Field, field_validator
from typing import Generator, List, Literal, Optional
from datetime import date
from urllib.parse import urlparse


PROFILE_DB = Path(
    os.getenv(
        "DATABASE_PATH",
        str(Path(__file__).resolve().parent / "profiles.db"),
    )
).expanduser()


@contextmanager
def profile_db_connection() -> Generator[sqlite3.Connection, None, None]:
    connection = sqlite3.connect(PROFILE_DB)
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def init_profile_db():
    with profile_db_connection() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS profiles (
                user_id TEXT PRIMARY KEY,
                full_name TEXT NOT NULL DEFAULT '',
                email TEXT NOT NULL DEFAULT '',
                education TEXT NOT NULL DEFAULT '',
                university TEXT NOT NULL DEFAULT '',
                graduation_year INTEGER,
                skills TEXT NOT NULL DEFAULT '[]',
                projects TEXT NOT NULL DEFAULT '[]',
                experience TEXT NOT NULL DEFAULT '',
                career_interests TEXT NOT NULL DEFAULT '[]',
                preferred_locations TEXT NOT NULL DEFAULT '[]',
                opportunity_types TEXT NOT NULL DEFAULT '[]',
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS applications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id TEXT NOT NULL,
                opportunity_id TEXT NOT NULL,
                job_title TEXT NOT NULL,
                employer_name TEXT NOT NULL DEFAULT '',
                job_url TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'Saved'
                    CHECK (status IN ('Saved', 'Applied', 'Interview', 'Offer', 'Rejected')),
                application_date TEXT,
                notes TEXT NOT NULL DEFAULT '',
                follow_up_date TEXT,
                application_deadline_date TEXT,
                deadline_source TEXT NOT NULL DEFAULT '',
                follow_up_completed INTEGER NOT NULL DEFAULT 0,
                application_deadline_completed INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE (user_id, opportunity_id)
            )
        """)
        application_columns = {
            row[1]
            for row in conn.execute("PRAGMA table_info(applications)").fetchall()
        }
        if "application_deadline_date" not in application_columns:
            conn.execute(
                "ALTER TABLE applications ADD COLUMN application_deadline_date TEXT"
            )
        if "deadline_source" not in application_columns:
            conn.execute(
                "ALTER TABLE applications ADD COLUMN deadline_source TEXT NOT NULL DEFAULT ''"
            )
        if "application_deadline_completed" not in application_columns:
            conn.execute(
                "ALTER TABLE applications ADD COLUMN application_deadline_completed INTEGER NOT NULL DEFAULT 0"
            )
        conn.execute("""
            CREATE INDEX IF NOT EXISTS applications_user_status_idx
            ON applications (user_id, status)
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS deadlines (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id TEXT NOT NULL,
                title TEXT NOT NULL,
                due_date TEXT,
                source TEXT NOT NULL DEFAULT '',
                completed INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE INDEX IF NOT EXISTS deadlines_user_due_idx
            ON deadlines (user_id, due_date)
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS roadmap_tasks (
                user_id TEXT NOT NULL,
                task_id TEXT NOT NULL,
                title TEXT NOT NULL,
                description TEXT NOT NULL DEFAULT '',
                category TEXT NOT NULL DEFAULT '',
                priority INTEGER NOT NULL DEFAULT 0,
                completed INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY (user_id, task_id)
            )
        """)


init_profile_db()


class StudentProfile(BaseModel):
    user_id: str = Field(min_length=1, max_length=100)
    full_name: str = Field(default="", max_length=200)
    email: str = Field(default="", max_length=254)
    education: str = Field(default="", max_length=300)
    university: str = Field(default="", max_length=300)
    graduation_year: Optional[int] = Field(default=None, ge=1900, le=2100)
    skills: List[str] = Field(default_factory=list, max_length=100)
    projects: List[str] = Field(default_factory=list, max_length=100)
    experience: str = Field(default="", max_length=5000)
    career_interests: List[str] = Field(default_factory=list, max_length=100)
    preferred_locations: List[str] = Field(default_factory=list, max_length=100)
    opportunity_types: List[str] = Field(default_factory=list, max_length=50)

    @field_validator("user_id")
    @classmethod
    def normalize_user_id(cls, value):
        value = value.strip()
        if not value:
            raise ValueError("User ID cannot be blank.")
        return value

    @field_validator(
        "skills",
        "projects",
        "career_interests",
        "preferred_locations",
        "opportunity_types",
    )
    @classmethod
    def normalize_profile_lists(cls, values):
        return [
            value.strip()[:300]
            for value in values
            if value.strip()
        ]

    @field_validator("email")
    @classmethod
    def validate_optional_email(cls, value):
        value = value.strip()
        if value and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value):
            raise ValueError("Enter a valid email address.")
        return value


def profile_to_dict(row):
    return {
        "user_id": row[0],
        "full_name": row[1],
        "email": row[2],
        "education": row[3],
        "university": row[4],
        "graduation_year": row[5],
        "skills": json.loads(row[6]),
        "projects": json.loads(row[7]),
        "experience": row[8],
        "career_interests": json.loads(row[9]),
        "preferred_locations": json.loads(row[10]),
        "opportunity_types": json.loads(row[11]),
        "updated_at": row[12],
    }


@app.post("/profile")
async def save_profile(profile: StudentProfile):
    with profile_db_connection() as conn:
        conn.execute("""
            INSERT INTO profiles (
                user_id, full_name, email, education, university,
                graduation_year, skills, projects, experience,
                career_interests, preferred_locations, opportunity_types,
                updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id) DO UPDATE SET
                full_name = excluded.full_name,
                email = excluded.email,
                education = excluded.education,
                university = excluded.university,
                graduation_year = excluded.graduation_year,
                skills = excluded.skills,
                projects = excluded.projects,
                experience = excluded.experience,
                career_interests = excluded.career_interests,
                preferred_locations = excluded.preferred_locations,
                opportunity_types = excluded.opportunity_types,
                updated_at = CURRENT_TIMESTAMP
        """, (
            profile.user_id,
            profile.full_name,
            profile.email,
            profile.education,
            profile.university,
            profile.graduation_year,
            json.dumps(profile.skills),
            json.dumps(profile.projects),
            profile.experience,
            json.dumps(profile.career_interests),
            json.dumps(profile.preferred_locations),
            json.dumps(profile.opportunity_types),
        ))

    return {
        "message": "Profile saved successfully",
        "profile": profile.model_dump()
    }


@app.get("/profile/{user_id}")
async def get_profile(user_id: str):
    with profile_db_connection() as conn:
        row = conn.execute("""
            SELECT user_id, full_name, email, education, university,
                   graduation_year, skills, projects, experience,
                   career_interests, preferred_locations,
                   opportunity_types, updated_at
            FROM profiles WHERE user_id = ?
        """, (user_id,)).fetchone()

    if row is None:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Profile not found")

    return profile_to_dict(row)


@app.delete("/profile/{user_id}")
async def delete_profile(user_id: str):
    with profile_db_connection() as conn:
        cursor = conn.execute(
            "DELETE FROM profiles WHERE user_id = ?", (user_id,)
        )

    if cursor.rowcount == 0:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Profile not found")

    return {"message": "Profile deleted successfully"}


class ApplicationInput(BaseModel):
    user_id: str = Field(min_length=1, max_length=100)
    opportunity_id: str = Field(min_length=1, max_length=2000)
    job_title: str = Field(min_length=1, max_length=300)
    employer_name: str = Field(default="", max_length=300)
    job_url: str = Field(default="", max_length=2000)
    status: Literal["Saved", "Applied", "Interview", "Offer", "Rejected"] = "Saved"
    application_date: Optional[date] = None
    notes: str = Field(default="", max_length=5000)
    follow_up_date: Optional[date] = None
    application_deadline_date: Optional[date] = None
    deadline_source: str = Field(default="", max_length=2000)

    @field_validator("user_id", "opportunity_id", "job_title")
    @classmethod
    def require_application_identity(cls, value):
        value = value.strip()
        if not value:
            raise ValueError("This field cannot be blank.")
        return value

    @field_validator("employer_name", "notes", "deadline_source")
    @classmethod
    def trim_application_text(cls, value):
        return value.strip()

    @field_validator("job_url")
    @classmethod
    def validate_job_url(cls, value):
        value = value.strip()
        if not value:
            return value
        parsed = urlparse(value)
        if parsed.scheme not in ("http", "https") or not parsed.netloc:
            raise ValueError("Job URL must use http or https.")
        return value


class DeadlineInput(BaseModel):
    user_id: str = Field(min_length=1, max_length=100)
    title: str = Field(min_length=1, max_length=300)
    due_date: Optional[date] = None
    source: str = Field(default="", max_length=1000)

    @field_validator("user_id", "title")
    @classmethod
    def require_deadline_identity(cls, value):
        value = value.strip()
        if not value:
            raise ValueError("This field cannot be blank.")
        return value

    @field_validator("source")
    @classmethod
    def trim_deadline_source(cls, value):
        return value.strip()


class DeadlineCompletion(BaseModel):
    user_id: str = Field(min_length=1, max_length=100)
    completed: bool


class RoadmapTaskInput(BaseModel):
    user_id: str = Field(min_length=1, max_length=100)
    task_id: str = Field(min_length=1, max_length=100)
    title: str = Field(min_length=1, max_length=300)
    description: str = Field(default="", max_length=2000)
    category: str = Field(default="", max_length=100)
    priority: int = Field(default=0, ge=0, le=100)
    completed: bool = False

    @field_validator("user_id", "task_id", "title")
    @classmethod
    def require_roadmap_identity(cls, value):
        value = value.strip()
        if not value:
            raise ValueError("This field cannot be blank.")
        return value


class RoadmapTaskCompletion(BaseModel):
    user_id: str = Field(min_length=1, max_length=100)
    completed: bool


class MatchRequest(BaseModel):
    profile: StudentProfile
    opportunity: dict


class MatchBatchRequest(BaseModel):
    profile: StudentProfile
    opportunities: List[dict]


@app.get("/health")
def health():
    return {"status": "ok"}


def application_from_row(row):
    return {
        "id": row[0],
        "user_id": row[1],
        "opportunity_id": row[2],
        "job_title": row[3],
        "employer_name": row[4],
        "job_url": row[5],
        "status": row[6],
        "application_date": row[7],
        "notes": row[8],
        "follow_up_date": row[9],
        "application_deadline_date": row[10],
        "deadline_source": row[11],
        "follow_up_completed": bool(row[12]),
        "application_deadline_completed": bool(row[13]),
        "created_at": row[14],
        "updated_at": row[15],
    }


APPLICATION_COLUMNS = """
    id, user_id, opportunity_id, job_title, employer_name, job_url,
    status, application_date, notes, follow_up_date, application_deadline_date,
    deadline_source, follow_up_completed, application_deadline_completed,
    created_at, updated_at
"""


@app.get("/applications/{user_id}")
def list_applications(user_id: str):
    with profile_db_connection() as conn:
        rows = conn.execute(
            f"""
            SELECT {APPLICATION_COLUMNS} FROM applications
            WHERE user_id = ?
            ORDER BY updated_at DESC, id DESC
            """,
            (user_id,),
        ).fetchall()

    return [application_from_row(row) for row in rows]


@app.post("/applications")
def save_application(application: ApplicationInput):
    values = application.model_dump()
    values["application_date"] = (
        values["application_date"].isoformat()
        if values["application_date"]
        else None
    )
    values["follow_up_date"] = (
        values["follow_up_date"].isoformat()
        if values["follow_up_date"]
        else None
    )
    values["application_deadline_date"] = (
        values["application_deadline_date"].isoformat()
        if values["application_deadline_date"]
        else None
    )

    with profile_db_connection() as conn:
        conn.execute(
            """
            INSERT INTO applications (
                user_id, opportunity_id, job_title, employer_name, job_url,
                status, application_date, notes, follow_up_date,
                application_deadline_date, deadline_source, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, opportunity_id) DO UPDATE SET
                job_title = excluded.job_title,
                employer_name = excluded.employer_name,
                job_url = excluded.job_url,
                status = excluded.status,
                application_date = excluded.application_date,
                notes = excluded.notes,
                follow_up_date = excluded.follow_up_date,
                application_deadline_date = excluded.application_deadline_date,
                deadline_source = excluded.deadline_source,
                application_deadline_completed = CASE
                    WHEN applications.application_deadline_date
                         IS NOT excluded.application_deadline_date
                    THEN 0 ELSE applications.application_deadline_completed
                END,
                updated_at = CURRENT_TIMESTAMP
            """,
            (
                values["user_id"],
                values["opportunity_id"],
                values["job_title"],
                values["employer_name"],
                values["job_url"],
                values["status"],
                values["application_date"],
                values["notes"],
                values["follow_up_date"],
                values["application_deadline_date"],
                values["deadline_source"],
            ),
        )
        row = conn.execute(
            f"""
            SELECT {APPLICATION_COLUMNS} FROM applications
            WHERE user_id = ? AND opportunity_id = ?
            """,
            (values["user_id"], values["opportunity_id"]),
        ).fetchone()

    return application_from_row(row)


@app.put("/applications/{application_id}")
def update_application(application_id: int, application: ApplicationInput):
    values = application.model_dump()
    values["application_date"] = (
        values["application_date"].isoformat()
        if values["application_date"]
        else None
    )
    values["follow_up_date"] = (
        values["follow_up_date"].isoformat()
        if values["follow_up_date"]
        else None
    )
    values["application_deadline_date"] = (
        values["application_deadline_date"].isoformat()
        if values["application_deadline_date"]
        else None
    )

    with profile_db_connection() as conn:
        cursor = conn.execute(
            """
            UPDATE applications
            SET opportunity_id = ?, job_title = ?, employer_name = ?,
                job_url = ?, status = ?, application_date = ?, notes = ?,
                follow_up_date = ?, application_deadline_date = ?,
                deadline_source = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND user_id = ?
            """,
            (
                values["opportunity_id"],
                values["job_title"],
                values["employer_name"],
                values["job_url"],
                values["status"],
                values["application_date"],
                values["notes"],
                values["follow_up_date"],
                values["application_deadline_date"],
                values["deadline_source"],
                application_id,
                values["user_id"],
            ),
        )
        row = conn.execute(
            f"SELECT {APPLICATION_COLUMNS} FROM applications WHERE id = ?",
            (application_id,),
        ).fetchone()

    if cursor.rowcount == 0 or row is None:
        raise HTTPException(status_code=404, detail="Application not found.")

    return application_from_row(row)


@app.delete("/applications/{application_id}")
def remove_application(application_id: int, user_id: str = Query(min_length=1)):
    with profile_db_connection() as conn:
        cursor = conn.execute(
            "DELETE FROM applications WHERE id = ? AND user_id = ?",
            (application_id, user_id),
        )

    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Application not found.")

    return {"message": "Application deleted."}


@app.get("/deadlines/{user_id}")
def list_deadlines(user_id: str):
    with profile_db_connection() as conn:
        deadlines = conn.execute(
            """
            SELECT id, title, due_date, source, completed, created_at
            FROM deadlines
            WHERE user_id = ?
            """,
            (user_id,),
        ).fetchall()
        application_actions = conn.execute(
            """
            SELECT id, job_title, job_url, follow_up_date,
                   follow_up_completed, created_at,
                   application_deadline_date, deadline_source,
                   application_deadline_completed
            FROM applications
            WHERE user_id = ?
              AND (follow_up_date IS NOT NULL OR application_deadline_date IS NOT NULL)
            """,
            (user_id,),
        ).fetchall()

    items = [
        {
            "id": str(row[0]),
            "title": row[1],
            "due_date": row[2],
            "source": row[3],
            "completed": bool(row[4]),
            "created_at": row[5],
            "kind": "deadline",
        }
        for row in deadlines
    ]
    for row in application_actions:
        if row[3]:
            items.append({
                "id": f"application-{row[0]}",
                "title": f"Follow up: {row[1]}",
                "due_date": row[3],
                "source": "Application follow-up",
                "completed": bool(row[4]),
                "created_at": row[5],
                "kind": "follow-up",
            })
        if row[6]:
            items.append({
                "id": f"application-deadline-{row[0]}",
                "title": f"Application closes: {row[1]}",
                "due_date": row[6],
                "source": row[7] or row[2] or "Job listing",
                "completed": bool(row[8]),
                "created_at": row[5],
                "kind": "application-deadline",
            })
    return sorted(items, key=lambda item: (item["due_date"] is None, item["due_date"] or ""))


@app.post("/deadlines")
def create_deadline(deadline: DeadlineInput):
    with profile_db_connection() as conn:
        cursor = conn.execute(
            """
            INSERT INTO deadlines (user_id, title, due_date, source)
            VALUES (?, ?, ?, ?)
            """,
            (
                deadline.user_id,
                deadline.title.strip(),
                deadline.due_date.isoformat() if deadline.due_date else None,
                deadline.source.strip(),
            ),
        )
        deadline_id = cursor.lastrowid
        row = conn.execute(
            """
            SELECT id, title, due_date, source, completed, created_at
            FROM deadlines WHERE id = ?
            """,
            (deadline_id,),
        ).fetchone()

    return {
        "id": str(row[0]),
        "title": row[1],
        "due_date": row[2],
        "source": row[3],
        "completed": bool(row[4]),
        "created_at": row[5],
        "kind": "deadline",
    }


@app.patch("/deadlines/{deadline_id}")
def set_deadline_completion(
    deadline_id: str,
    update: DeadlineCompletion,
):
    if deadline_id.startswith("application-deadline-"):
        try:
            application_id = int(deadline_id.removeprefix("application-deadline-"))
        except ValueError as error:
            raise HTTPException(
                status_code=404,
                detail="Deadline not found.",
            ) from error

        with profile_db_connection() as conn:
            cursor = conn.execute(
                """
                UPDATE applications
                SET application_deadline_completed = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND user_id = ?
                  AND application_deadline_date IS NOT NULL
                """,
                (int(update.completed), application_id, update.user_id),
            )
    elif deadline_id.startswith("application-"):
        try:
            application_id = int(deadline_id.removeprefix("application-"))
        except ValueError as error:
            raise HTTPException(
                status_code=404,
                detail="Deadline not found.",
            ) from error

        with profile_db_connection() as conn:
            cursor = conn.execute(
                """
                UPDATE applications
                SET follow_up_completed = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND user_id = ? AND follow_up_date IS NOT NULL
                """,
                (int(update.completed), application_id, update.user_id),
            )
    else:
        try:
            numeric_id = int(deadline_id)
        except ValueError as error:
            raise HTTPException(
                status_code=404,
                detail="Deadline not found.",
            ) from error

        with profile_db_connection() as conn:
            cursor = conn.execute(
                """
                UPDATE deadlines
                SET completed = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND user_id = ?
                """,
                (int(update.completed), numeric_id, update.user_id),
            )

    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Deadline not found.")

    return {"id": deadline_id, "completed": update.completed}


@app.delete("/deadlines/{deadline_id}")
def remove_deadline(deadline_id: str, user_id: str = Query(min_length=1)):
    if deadline_id.startswith("application-deadline-"):
        try:
            application_id = int(deadline_id.removeprefix("application-deadline-"))
        except ValueError as error:
            raise HTTPException(
                status_code=404,
                detail="Deadline not found.",
            ) from error

        with profile_db_connection() as conn:
            cursor = conn.execute(
                """
                UPDATE applications
                SET application_deadline_date = NULL,
                    application_deadline_completed = 0,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND user_id = ?
                """,
                (application_id, user_id),
            )
    elif deadline_id.startswith("application-"):
        try:
            application_id = int(deadline_id.removeprefix("application-"))
        except ValueError as error:
            raise HTTPException(
                status_code=404,
                detail="Deadline not found.",
            ) from error

        with profile_db_connection() as conn:
            cursor = conn.execute(
                """
                UPDATE applications
                SET follow_up_date = NULL, follow_up_completed = 0,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND user_id = ?
                """,
                (application_id, user_id),
            )
    else:
        try:
            numeric_id = int(deadline_id)
        except ValueError as error:
            raise HTTPException(
                status_code=404,
                detail="Deadline not found.",
            ) from error

        with profile_db_connection() as conn:
            cursor = conn.execute(
                "DELETE FROM deadlines WHERE id = ? AND user_id = ?",
                (numeric_id, user_id),
            )

    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Deadline not found.")

    return {"message": "Deadline removed."}


@app.get("/roadmap/{user_id}")
def list_roadmap(user_id: str):
    with profile_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT user_id, task_id, title, description, category,
                   priority, completed, updated_at
            FROM roadmap_tasks
            WHERE user_id = ?
            ORDER BY priority, task_id
            """,
            (user_id,),
        ).fetchall()

    return [
        {
            "user_id": row[0],
            "task_id": row[1],
            "title": row[2],
            "description": row[3],
            "category": row[4],
            "priority": row[5],
            "completed": bool(row[6]),
            "updated_at": row[7],
        }
        for row in rows
    ]


@app.post("/roadmap")
def save_roadmap_task(task: RoadmapTaskInput):
    with profile_db_connection() as conn:
        conn.execute(
            """
            INSERT INTO roadmap_tasks (
                user_id, task_id, title, description, category,
                priority, completed, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, task_id) DO UPDATE SET
                title = excluded.title,
                description = excluded.description,
                category = excluded.category,
                priority = excluded.priority,
                updated_at = CURRENT_TIMESTAMP
            """,
            (
                task.user_id,
                task.task_id,
                task.title.strip(),
                task.description.strip(),
                task.category.strip(),
                task.priority,
                int(task.completed),
            ),
        )

    return task.model_dump()


@app.patch("/roadmap/{task_id}")
def set_roadmap_task_completion(
    task_id: str,
    update: RoadmapTaskCompletion,
):
    with profile_db_connection() as conn:
        cursor = conn.execute(
            """
            UPDATE roadmap_tasks
            SET completed = ?, updated_at = CURRENT_TIMESTAMP
            WHERE task_id = ? AND user_id = ?
            """,
            (int(update.completed), task_id, update.user_id),
        )

    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Roadmap task not found.")

    return {"task_id": task_id, "completed": update.completed}


@app.delete("/roadmap/{task_id}")
def remove_roadmap_task(task_id: str, user_id: str = Query(min_length=1)):
    with profile_db_connection() as conn:
        cursor = conn.execute(
            "DELETE FROM roadmap_tasks WHERE task_id = ? AND user_id = ?",
            (task_id, user_id),
        )

    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Roadmap task not found.")

    return {"message": "Roadmap task removed."}


def calculate_match(profile: StudentProfile, opportunity: dict) -> dict:
    title = str(opportunity.get("job_title") or "")
    description = str(opportunity.get("job_description") or "")
    job_text = f"{title}\n{description}".casefold()
    dimensions = []
    reasons = []
    profile_skills = {skill.strip().casefold() for skill in profile.skills}
    job_skills = find_skills(description)
    matching_skills = [
        skill for skill in job_skills if skill.casefold() in profile_skills
    ]
    missing_skills = [
        skill for skill in job_skills if skill.casefold() not in profile_skills
    ]

    if job_skills:
        dimensions.append((50, len(matching_skills) / len(job_skills)))
        if matching_skills:
            reasons.append("Matching skills: " + ", ".join(matching_skills))
        if missing_skills:
            reasons.append("Skills to develop: " + ", ".join(missing_skills))
    else:
        reasons.append("The listing does not provide detectable skill requirements.")

    if profile.career_interests and job_text.strip():
        matching_interests = [
            interest for interest in profile.career_interests
            if interest.strip().casefold() in job_text
        ]
        dimensions.append(
            (20, 1.0 if matching_interests else 0.0)
        )
        if matching_interests:
            reasons.append(
                "Related career interests: " + ", ".join(matching_interests)
            )
        else:
            reasons.append("No saved career interest was found in the listing text.")

    education_requirements = {
        "bachelor": ("bachelor", "b.tech", "b.e.", "b.sc", "undergraduate"),
        "master": ("master", "m.tech", "m.e.", "m.sc", "postgraduate"),
        "doctorate": ("phd", "doctorate"),
    }
    required_levels = [
        level for level, terms in education_requirements.items()
        if any(term in job_text for term in terms)
    ]
    if required_levels:
        education = profile.education.casefold()
        education_matches = any(
            term in education
            for level in required_levels
            for term in education_requirements[level]
        )
        dimensions.append((10, 1.0 if education_matches else 0.0))
        reasons.append(
            "Education text matches the listed degree level."
            if education_matches
            else "The listed degree requirement is not confirmed by the saved education."
        )

    preferred_locations = [
        location.strip().casefold()
        for location in profile.preferred_locations
        if location.strip()
    ]
    job_location = " ".join(
        str(opportunity.get(key) or "")
        for key in ("job_city", "job_state", "job_country")
    ).casefold()
    is_remote = bool(opportunity.get("job_is_remote"))
    if preferred_locations and (job_location.strip() or is_remote):
        location_match = any(
            location in job_location
            or (is_remote and location == "remote")
            for location in preferred_locations
        )
        dimensions.append((10, 1.0 if location_match else 0.0))
        reasons.append(
            "The location matches a saved preference."
            if location_match
            else "The listing location does not match saved preferences."
        )

    preferred_types = [
        re.sub(r"[^a-z]", "", value.casefold())
        for value in profile.opportunity_types
    ]
    job_type = re.sub(
        r"[^a-z]",
        "",
        str(opportunity.get("job_employment_type") or "").casefold(),
    )
    if preferred_types and job_type:
        type_match = any(
            preferred in job_type or job_type in preferred
            for preferred in preferred_types
            if preferred
        )
        dimensions.append((10, 1.0 if type_match else 0.0))
        reasons.append(
            "The opportunity type matches a saved preference."
            if type_match
            else "The opportunity type differs from saved preferences."
        )

    total_weight = sum(weight for weight, _ in dimensions)
    score = (
        round(100 * sum(weight * value for weight, value in dimensions) / total_weight)
        if total_weight
        else None
    )

    return {
        "score": score,
        "reasons": reasons,
        "matching_skills": matching_skills,
        "missing_skills": missing_skills,
        "scored_dimensions": len(dimensions),
        "note": (
            "This score is a transparent heuristic, not a prediction of selection."
        ),
    }


@app.post("/matching/score")
def score_opportunity(request: MatchRequest):
    return calculate_match(request.profile, request.opportunity)


@app.post("/matching/scores")
def score_opportunities(request: MatchBatchRequest):
    return [
        {
            "opportunity_id": (
                opportunity.get("job_id")
                or opportunity.get("job_apply_link")
                or (
                    f"{opportunity.get('job_title') or 'job'}-"
                    f"{opportunity.get('employer_name') or index}"
                )
            ),
            **calculate_match(request.profile, opportunity),
        }
        for index, opportunity in enumerate(request.opportunities)
    ]
